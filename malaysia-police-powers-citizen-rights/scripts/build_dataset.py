#!/usr/bin/env python3
"""Build the Malaysian Police Powers & Citizen Rights dataset (CSV + XLSX).

Usage:  python3 scripts/build_dataset.py
Writes to ../data/ relative to this file.
"""
import csv
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from rows.common import COLUMNS  # noqa: E402
from rows import (  # noqa: E402
    cat_a_street, cat_b_arrest, cat_cd_detention_remand, cat_e_body_search,
    cat_f_property, cat_g_digital, cat_hi_vehicle_drugs, cat_j_foreign,
)
import companion_tables as ct  # noqa: E402

DATA_DIR = Path(__file__).resolve().parent.parent / "data"

ALL_ROWS = (
    cat_a_street.ROWS + cat_b_arrest.ROWS + cat_cd_detention_remand.ROWS
    + cat_e_body_search.ROWS + cat_f_property.ROWS + cat_g_digital.ROWS
    + cat_hi_vehicle_drugs.ROWS + cat_j_foreign.ROWS
)


def validate(rows):
    seen = set()
    for r in rows:
        assert set(r) == set(COLUMNS), f"column mismatch in {r['Scenario']}"
        num = int(r["Scenario"].split(".")[0])
        assert num not in seen, f"duplicate scenario number {num}"
        seen.add(num)
        assert r["Confidence_Level"] in {"HIGH", "MEDIUM", "LOW"}, r["Scenario"]
        for c in ("Police_Power", "Required_Legal_Trigger", "Citizen_Questions",
                  "How_To_Verify_Legality", "Red_Flags", "Exact_Citation"):
            assert r[c].strip(), f"empty {c} in {r['Scenario']}"
        if r["Category"].startswith("G"):
            assert "PHONE SEARCH FRAMEWORK" in r["Research_Notes"], r["Scenario"]
    assert seen == set(range(1, 70)), sorted(set(range(1, 70)) - seen)
    return rows


def write_csv(path, columns, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL, lineterminator="\n")
        w.writerow(columns)
        for r in rows:
            w.writerow([r[c] if isinstance(r, dict) else r[i] for i, c in enumerate(columns)])


def write_xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for name, columns, rows, widths in sheets:
        ws = wb.create_sheet(title=name[:31])
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([r[c] if isinstance(r, dict) else r[i] for i, c in enumerate(columns)])
        for i, c in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 28)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2" if name == "Scenarios" else "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    rows = validate(ALL_ROWS)

    write_csv(DATA_DIR / "police_powers_citizen_rights.csv", COLUMNS, rows)
    write_csv(DATA_DIR / "research_gaps.csv", ct.GAPS_COLUMNS, ct.RESEARCH_GAPS)
    write_csv(DATA_DIR / "top20_rights.csv", ct.RIGHTS_COLUMNS, ct.TOP20_RIGHTS)
    write_csv(DATA_DIR / "top20_police_powers.csv", ct.POWERS_COLUMNS, ct.TOP20_POWERS)
    write_csv(DATA_DIR / "sources_register.csv", ct.SOURCES_COLUMNS, ct.SOURCES)
    write_csv(DATA_DIR / "data_dictionary.csv", ct.DICT_COLUMNS, ct.DATA_DICTIONARY)

    scen_widths = {"Category": 18, "Scenario": 30, "Section": 30, "Exact_Citation": 60,
                   "Confidence_Level": 12, "Need_Warrant": 16, "Need_Arrest": 16,
                   "Need_Reasonable_Suspicion": 16, "Need_Judicial_Approval": 16}
    write_xlsx(DATA_DIR / "police_powers_citizen_rights.xlsx", [
        ("Scenarios", COLUMNS, rows, scen_widths),
        ("Research_Gaps", ct.GAPS_COLUMNS, ct.RESEARCH_GAPS, {"Issue": 30, "Comments": 50}),
        ("Top20_Rights", ct.RIGHTS_COLUMNS, ct.TOP20_RIGHTS, {"Rank": 6, "Right": 50, "Why_It_Matters": 50}),
        ("Top20_Powers", ct.POWERS_COLUMNS, ct.TOP20_POWERS, {"Rank": 6, "Police_Power": 50, "Key_Limits": 50}),
        ("Sources_Register", ct.SOURCES_COLUMNS, ct.SOURCES, {"Title": 45, "Coverage_Notes": 60}),
        ("Data_Dictionary", ct.DICT_COLUMNS, ct.DATA_DICTIONARY, {"Definition": 60, "Allowed_Values_Or_Format": 50}),
    ])

    conf = {k: sum(1 for r in rows if r["Confidence_Level"] == k) for k in ("HIGH", "MEDIUM", "LOW")}
    print(f"rows={len(rows)} confidence={conf}")
    print(f"gaps={len(ct.RESEARCH_GAPS)} rights={len(ct.TOP20_RIGHTS)} powers={len(ct.TOP20_POWERS)} sources={len(ct.SOURCES)}")


if __name__ == "__main__":
    main()
