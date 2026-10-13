#!/usr/bin/env python3
"""Annotate the reviewer's workbook with a QA verdict on every reviewer entry.

Adds two columns to each reviewed tab: 'QA verdict' and 'QA notes'.
QA verdict values: OK | OK-NOTE | NOT OK | UNVERIFIABLE
ID and Hash columns are never touched.

Usage: python3 qa_annotate.py <reviewed.xlsx> <out.xlsx>
"""
import sys
from collections import Counter
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SECOND_PERSON = (
    "OK-NOTE",
    "Amend is defensible only if a third-person house style is adopted. The Read me asks for "
    "'information rather than instruction'; 'You may have the right...' is information, not instruction, "
    "so this verdict rests on a standard the Read me does not state. The same standard was not applied to the "
    "Rules tab (31 of 31 'The app shows' texts are second person and all Approved). Decide the style once and apply it to every tab.",
)

ASSERTIONS = {
    # second-person Amends
    **{k: SECOND_PERSON for k in [
        "ast.handcuffing", "ast.no_coercion_arrest", "ast.same_sex_search_right", "ast.lawyer_access_detention",
        "ast.magistrate_24h_constrains", "ast.remand_representations", "ast.officer_identification_right",
        "ast.vehicle_search_grounds_right", "ast.vehicle_search_right_trigger", "ast.leave_when_not_arrested",
        "ast.what_is_arrest", "ast.identity_card_gap_street", "ast.section112_self_incrimination",
        "ast.section112_self_incrimination_trigger", "ast.witness_not_arrested", "ast.witness_not_arrested_trigger",
    ]},
    "ast.habeas_corpus_sosma": ("NOT OK",
        "Holding a draft back is reasonable, but the stated reason is legally overstated. SOSMA s.4(10) and s.5(3) say those sections "
        "'shall have effect notwithstanding anything inconsistent with' Art 5; they do not say Art 5(2) is excluded, and s.4(10) "
        "only protects detention that is within s.4. Habeas corpus for detention outside s.4 (beyond 28 days, no Superintendent "
        "authorisation, arrest solely for political belief contrary to s.4(3)) remains arguable. Change to 'Amend - hold for counsel; "
        "Requires Case Law Review' and correct the reasoning. Do not record that SOSMA 'excludes Art 5 rights except s.4(2)'."),
    "ast.sosma_extension_rank": ("OK-NOTE",
        "Amend note is accurate to s.4(11). But the Earlier note asked the reviewer to 'confirm the current renewal before approval' and "
        "the reviewer did not record whether s.4(5) is currently in force (last parliamentary renewal date). That check is still open."),
    "ast.lethal_force_limit": ("OK-NOTE",
        "Approve is correct and '4 July 2023' matches the Penal Code reprint supplied ('Act 846 ... came into operation on 4 July 2023'). "
        "It contradicts the reviewer's own Sources & currency entry giving Act 846 commencement as 16/6/2023. Fix the Sources tab."),
    "ast.warrant_of_arrest_route": ("OK", "Correct: text placeholder must be filled before publication."),
    "ast.witness_order_disobedience_gap": ("OK-NOTE",
        "Amend is defensible and the Penal Code s.174 pointer is useful. The note calls the row a 'draft'; its Status is 'published'. "
        "Reword so the verdict ties to the right status."),
    "ast.phone_content_uncertain": ("NOT OK",
        "Approved 'verified against CMA s.249' using the 2006 reprint text, while the reviewer's own Sources tab records that the "
        "Communications and Multimedia (Amendment) Act 2025 (A1743) commenced 11/2/2025. Whether A1743 amended s.249 was the open "
        "question in the currency note and it was not answered. Cannot be Approved as current until that is confirmed."),
    "ast.redbook_4_1_grounds": ("OK-NOTE",
        "Boilerplate says 'verified against cited statutory text (Red Book para 4.1 ...)'. The Red Book is guidance, not statute. "
        "Harmless here, but the same boilerplate appears on every Red Book row and shows the notes were templated rather than written per row."),
    "ast.approach_general_duty": ("OK-NOTE",
        "Police Act s.20(3) lists duties; it contains no words about approaching or speaking to a person. The claim is a fair inference, "
        "but the note should say 'inferred from the general duties in s.20(3)', not 'verified against'."),
    "ast.remand_rearrest_gap": ("OK-NOTE",
        "Approve is fine for the claim text, but the row ID says 're-arrest' while the claim is about the 'well founded' standard. "
        "The reviewer did not flag the mismatch; the content team should check the ID/claim pairing."),
    "ast.to_station_limit": ("OK", "Correct, and the reviewer correctly relied on the Earlier correction (Magistrate, not 'nearest station')."),
    "ast.no_arrest_licence_and_name_limit": ("OK-NOTE",
        "Approve is correct on the 2019 text. The reviewer notes A1794 (2026) changes the s.112(1)(a) list but recorded it as 'not yet in force' "
        "in the Sources tab without an LOM reference; keep this row on a watch list."),
}

ENTITIES = {
    **{k: SECOND_PERSON for k in ["rgt.lawyer_access", "rgt.family_notification", "rgt.self_incrimination", "rgt.vehicle_search_grounds"]},
    "pwr.arrest_with_warrant": ("OK", "Correct: draft needs the ss.38-47 text."),
}
ENTITY_DEFAULT = ("OK-NOTE",
    "Verdict acceptable. Note is the same boilerplate on every row of this type and does not identify what was checked "
    "(e.g. which research rows or provisions). Acceptable for titles, which are short, but it is not an audit trail.")

PROVISIONS = {
    "prv.cpc_s15": ("NOT OK",
        "The note says the consolidated text was 'verified against ... criminal procedure code reprint 2022.pdf p.34'. The 2022 reprint "
        "cannot contain words inserted by Act 846 in 2023, and the note then copies the Earlier note's 'verify before approval' while "
        "approving. The inserted words must be checked against Act 846 s.61 (listed as supplied in the Sources tab) and the note must say so."),
    "prv.cpc_s116": ("NOT OK",
        "Approved as 'verified against ... p.52'. Section 116 is on printed page 83 of the 2022 reprint (s.111-112 are on p.81, s.116B on p.85, "
        "as the same tab records). Page 52 holds ss.51-52. The page check was not done."),
    "prv.cpc_s116a": ("NOT OK",
        "Approved as 'verified against ... p.57'. Section 116A is on printed page 84 (p.57 holds s.62). The page check was not done."),
    "prv.cca_s5": ("NOT OK",
        "Approved as 'verified against ... p.3'. CCA s.5 is on printed page 10 of the 2011 reprint (p.3 is the arrangement of sections)."),
    "prv.cca_s11": ("NOT OK",
        "Approved as 'verified against ... p.9'. CCA s.11 is on printed page 14."),
    "prv.rta_s78": ("NOT OK",
        "Approved as 'verified against ... p.57'. RTA s.78 is on printed pages 128-129 of the 2019 reprint."),
    "prv.cpc_s28": ("NOT OK",
        "Approved as 'verified against ... p.39'. Section 28 begins on printed page 41 of the 2022 reprint (p.39 holds s.24). "
        "Transcription is accurate; the page check was not done."),
    "prv.cpc_ch5_warrant": ("OK", "Correct."),
    "prv.cma_s249": ("NOT OK",
        "Column is 'Statutory text (latest version)'. The text is the 2006 reprint; the reviewer's Sources tab records A1743 (2025) as commenced "
        "11/2/2025. The reviewer must state whether A1743 touched s.249 before approving this as the latest version. "
        "(Text wording is consistent with s.249 as previously enacted; the issue is currency, not transcription.)"),
    "prv.rta_s55": ("OK-NOTE", "Text matches the 2019 reprint. The reviewer records a later RTA reprint (15/10/2023) in the Sources tab; confirm s.55 unchanged."),
    "prv.rta_s58": ("OK-NOTE", "Text matches the 2019 reprint. Confirm s.58 unchanged in the later reprint the reviewer cites."),
    "prv.rta_s112": ("OK-NOTE", "Text matches the 2019 reprint. Reviewer elsewhere notes A1794 (2026) amends s.112(1)(a); keep on watch list."),
    "prv.rta_s123": ("OK-NOTE", "Text matches the 2019 reprint. Confirm unchanged in the later reprint the reviewer cites."),
    "prv.sosma_s4": ("OK", "Text matches the 2015 reprint; A1736 (2024) deletes s.19 only."),
    "prv.sosma_s5": ("OK", "Text matches the 2015 reprint; A1736 (2024) deletes s.19 only."),
}
PROVISION_DEFAULT = ("OK", "Text and page reference consistent with the supplied reprint.")

ROUTES = {
    "Arrest under a warrant": ("NOT OK",
        "Note says 'verified against Criminal Procedure Code ss.38-47'. The pack contains no text for ss.38-47 (Provisions row prv.cpc_ch5_warrant "
        "is '[statutory text not yet added]' and the reviewer marked it Amend). The route cannot have been verified against text that is absent."),
    "Approach and speak to a person": ("OK-NOTE",
        "Police Act s.20(3) does not mention approaching or speaking to anyone; the route is an inference from the general duties. "
        "Say 'inferred from', not 'verified against'."),
    "Arrest and detention for security offences (SOSMA)": ("OK", "Correct, including the s.4(11) renewal caveat."),
}
ROUTE_DEFAULT = ("OK", "Route, threshold and citation match the supplied text. Note is formulaic but accurate.")

CALCULATOR = {
    "24-hour production": ("OK-NOTE", "Substance correct (FC Art 5(4); CPC s.28(3)). The bracketed citations [1][2] point to no reference list; remove them or add the list."),
    "Remand caps: under_14_years": ("OK-NOTE", "Substance correct (CPC s.117(2)(a), (3)-(5)). Dangling [n] citations."),
    "Remand caps: death_or_14_years_or_more": ("OK-NOTE", "Substance correct (CPC s.117(2)(b)). Dangling [n] citations."),
    "Regime: security": ("OK-NOTE", "Substance correct (SOSMA s.4(5), (10), (11)). Does not say whether s.4(5) is currently renewed. Dangling [n] citations."),
    "Regime: immigration": ("OK-NOTE", "Substance correct (FC Art 5(4) second proviso). Dangling [n] citations."),
    "Regime: drugs": ("OK-NOTE", "Substance correct (DDA s.31(1), s.37B; CPC s.117). Dangling [n] citations."),
}
CALC_OPEN = {
    "How remand days are counted (calendar days or 24-hour periods).": ("NOT OK",
        "The question was not answered. The reply restates the 24-hour rule and the s.117(2) caps; it never says whether a '4 day' order "
        "means four calendar days or 96 hours. CPC s.117 is silent; the answer needs the Interpretation Acts 1948 and 1967 (computation of time) "
        "or case law, with a citation."),
    "How a part-day is treated.": ("NOT OK",
        "The reply asserts 'a part of a day spent in custody is treated as a full day under statutory interpretation and legal practice [3][4]' "
        "with no authority. Nothing in CPC s.117 or the supplied texts says this. Unsupported; needs a cited source or 'unsettled'."),
    "Whether counting begins from the arrest or from the remand order.": ("NOT OK",
        "Answered for the 24-hour clock (from arrest, correct). Not answered for remand: the question is whether a remand order's days run from "
        "the order or from the arrest. s.117(3)-(4) only require prior detention to be disclosed and 'taken into consideration'; the reply "
        "stretches that into 'counts as part of the overall detention period evaluated by the court' without authority."),
    "Drugs regime: confirm the ordinary s.117 caps apply to DDA arrests by police (DDA s.31B customs arrests differ; not modelled).": ("OK-NOTE",
        "Substance correct (DDA s.31(1) seizable; s.31B(3) preserves CPC s.117 for police investigations). Dangling [n] citations."),
}

SOURCES = {
    "Federal Constitution": ("NOT OK", "Column asks 'Amendments checked up to (date)'. 'Reprint date - confirmed' gives no date and does not say whether any amendment after 30 Nov 2025 was found."),
    "Criminal Procedure Code": ("NOT OK", "No date given. The reviewer lists CPC amending Acts A1682 (2023) and A1722 (2024) below, both after the 1 Oct 2022 reprint, but does not say which cited sections they touch. That is the question this column exists to answer."),
    "Police Act 1967": ("NOT OK", "No date given. State the date LOM was checked and whether any amendment after 10 May 2024 exists."),
    "Road Transport Act 1987": ("NOT OK", "A later reprint date (15/10/2023) is asserted without an LOM reference, and the reviewer does not say which of ss.55, 58, 78, 112, 123 changed between 1 Apr 2019 and that reprint. The currency note asked exactly that."),
    "Dangerous Drugs Act 1952": ("NOT OK", "No date given."),
    "Security Offences (Special Measures) Act 2012": ("NOT OK", "No date given, and 'Act 747 NOT the one declared in the currency note' is unintelligible: Act 747 is SOSMA and the currency note names amending Act A1736. Clarify what was meant."),
    "Computer Crimes Act 1997": ("NOT OK", "No date given."),
    "Evidence Act 1950": ("NOT OK", "No date given. The reviewer lists Evidence (Amendment) Act 2024 (A1732) below, after the 2017 reprint, but does not say whether ss.24-26 are affected."),
    "Penal Code": ("NOT OK", "No date given. Reviewer lists A1691, A1721 and A1750 below, all after the 4 Jul 2023 reprint, without saying whether ss.330, 339-342, 348 are affected."),
    "Malaysian Anti-Corruption Commission Act 2009": ("UNVERIFIABLE", "'Reprint date: 1/1/2022' cannot be checked against the supplied sources (the supplied MACC text is the 2019 reprint). Give the LOM reference."),
    "Communications and Multimedia Act 1998": ("NOT OK", "'latest Oct 2026' is wrong on the material supplied: the October 2026 amending Act in the project (A1799, Royal Assent 25 Sep 2026) amends the MCMC Act 1998 (Act 589), not the CMA (Act 588). If a separate CMA amendment of Oct 2026 exists, cite it; otherwise delete."),
    "prv.cpc_s15": ("NOT OK", "Commencement of Act 846 given as 16/6/2023. The Penal Code reprint supplied states Act 846 'came into operation on 4 July 2023'; 16 June 2023 is the Gazette publication date. 'Repritn' typo."),
    "prv.cma_s249": ("NOT OK", "States A1743 commenced 11/2/2025 but the 'In force' column for s.249 still reads 'not confirmed' and the reviewer does not say whether s.249 itself was amended or commenced. The question asked was specifically about s.249."),
    "prv.rta_s112": ("UNVERIFIABLE", "'Reprint date 15/09/2026, not yet in force' for A1794 cannot be checked against supplied sources; give the LOM reference. If correct, the row is consistent with the Assertions tab."),
    "Abolition of Mandatory Death Penalty Act 2023 (Act 846)": ("NOT OK", "Same error: commencement 16/6/2023 contradicts the 4 July 2023 date in the Penal Code reprint supplied."),
    "Criminal Procedure Code (Amendment) Act 2023 (Act A1682)": ("UNVERIFIABLE", "Dates cannot be checked against supplied sources. More important: the reviewer does not say which CPC sections A1682 amends and whether any cited section is affected."),
    "Criminal Procedure Code (Amendment) Act 2024 (Act A1722)": ("UNVERIFIABLE", "As above: give the LOM reference and state whether any cited CPC section is affected."),
    "Penal Code (Amendment) Act 2023 (Act A1691)": ("UNVERIFIABLE", "Commencement 10/9/2025 cannot be checked; state whether ss.330, 339-342, 348 are affected."),
    "Penal Code (Amendment) Act 2024 (Act A1721)": ("UNVERIFIABLE", "As above."),
    "Penal Code (Amendment) Act 2025 (Act A1750)": ("UNVERIFIABLE", "As above."),
    "Security Offences (Special Measures) (Amendment) Act 2024 (Act A1736)": ("OK-NOTE", "Gazette date 1 Nov 2024 matches the supplied A1736. Commencement 03/03/2025 is identical to the entry for the Evidence (Amendment) Act 2024 below; confirm this is not a copy-paste. 'Reprint date' is the Gazette publication date for an amending Act; relabel."),
    "Evidence (Amendment) Act 2024 (Act A1732)": ("UNVERIFIABLE", "Dates identical to the SOSMA amendment row above (01/11/2024 and 03/03/2025); confirm independently. State whether ss.24-26 are affected."),
    "Communications and Multimedia (Amendment) Act 2025 (Act A1743)": ("NOT OK", "The currency note asked to confirm the commencement date for s.249 specifically (the Act may commence in parts). A single commencement date was given with no statement about s.249."),
    "Road Transport (Amendment) Act 2026 (Act A1794)": ("UNVERIFIABLE", "Cannot be checked against supplied sources; give the LOM reference."),
}
SOURCES_CHECKED_BY = ("OK-NOTE", "Name and method are given; add the date of the check. 'verifying' (present tense) reads as work in progress; write 'checked on <date>'.")


def style_header(cell):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="7B2D26")
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def add_cols(ws, key_col, mapping, default, key_fn=None):
    hdr = [c.value for c in ws[1]]
    c1, c2 = len(hdr) + 1, len(hdr) + 2
    ws.cell(row=1, column=c1, value="QA verdict"); style_header(ws.cell(row=1, column=c1))
    ws.cell(row=1, column=c2, value="QA notes"); style_header(ws.cell(row=1, column=c2))
    ws.column_dimensions[ws.cell(row=1, column=c1).column_letter].width = 14
    ws.column_dimensions[ws.cell(row=1, column=c2).column_letter].width = 80
    ki = hdr.index(key_col)
    counts = Counter()
    for r in range(2, ws.max_row + 1):
        key = ws.cell(row=r, column=ki + 1).value
        if key is None:
            continue
        key = key_fn(key) if key_fn else key
        verdict, note = mapping.get(key, default)
        ws.cell(row=r, column=c1, value=verdict)
        ws.cell(row=r, column=c2, value=note)
        for c in (c1, c2):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        counts[verdict] += 1
    return counts


def main(src, out):
    wb = load_workbook(src)
    summary = {}
    summary["Assertions"] = add_cols(wb["Assertions"], "ID", ASSERTIONS, ("OK", "Verdict and note acceptable; claim matches the cited text in the supplied reprint."))
    summary["Rules"] = add_cols(wb["Rules"], "ID", {}, ("OK-NOTE", "Verdict acceptable; text matches the relied-on claims. All 31 'The app shows' texts are second person and were Approved, which is inconsistent with the 20 second-person Amends on the Assertions and Entities tabs."))
    summary["Entities"] = add_cols(wb["Entities"], "ID", ENTITIES, ENTITY_DEFAULT)
    summary["Provisions"] = add_cols(wb["Provisions"], "ID", PROVISIONS, PROVISION_DEFAULT)
    summary["Routes"] = add_cols(wb["Routes"], "Power", ROUTES, ROUTE_DEFAULT)

    # Calculator: key on Item, except OPEN QUESTION rows keyed on Value
    ws = wb["Calculator"]
    hdr = [c.value for c in ws[1]]
    c1, c2 = len(hdr) + 1, len(hdr) + 2
    ws.cell(row=1, column=c1, value="QA verdict"); style_header(ws.cell(row=1, column=c1))
    ws.cell(row=1, column=c2, value="QA notes"); style_header(ws.cell(row=1, column=c2))
    ws.column_dimensions[ws.cell(row=1, column=c2).column_letter].width = 80
    cnt = Counter()
    for r in range(2, ws.max_row + 1):
        item, value = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        if item is None:
            continue
        verdict, note = (CALC_OPEN.get(value) if item == "OPEN QUESTION" else CALCULATOR.get(item)) or ("OK", "")
        ws.cell(row=r, column=c1, value=verdict); ws.cell(row=r, column=c2, value=note)
        for c in (c1, c2):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        cnt[verdict] += 1
    summary["Calculator"] = cnt

    # Sources & currency: key on first column; two blocks have sub-headers
    ws = wb["Sources & currency"]
    c1, c2 = 7, 8
    ws.cell(row=1, column=c1, value="QA verdict (col E)"); style_header(ws.cell(row=1, column=c1))
    ws.cell(row=1, column=c2, value="QA notes"); style_header(ws.cell(row=1, column=c2))
    ws.column_dimensions["H"].width = 80
    cnt = Counter()
    for r in range(2, ws.max_row + 1):
        key = ws.cell(row=r, column=1).value
        if key in (None, "Cited provisions named by a later amending Act", "Provision", "Amending instruments supplied"):
            continue
        verdict, note = SOURCES.get(key, ("OK", ""))
        ws.cell(row=r, column=c1, value=verdict); ws.cell(row=r, column=c2, value=note)
        for c in (c1, c2):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        cnt[verdict] += 1
    ws.cell(row=ws.max_row + 2, column=1, value="QA on 'Checked by' column (all rows)")
    ws.cell(row=ws.max_row, column=c1, value=SOURCES_CHECKED_BY[0]); ws.cell(row=ws.max_row, column=c2, value=SOURCES_CHECKED_BY[1])
    summary["Sources & currency"] = cnt

    # Summary sheet
    ss = wb.create_sheet("QA summary", 0)
    ss.append(["Tab", "OK", "OK-NOTE", "NOT OK", "UNVERIFIABLE", "Rows checked"])
    for c in ss[1]:
        style_header(c)
    for tab, cnt in summary.items():
        ss.append([tab, cnt.get("OK", 0), cnt.get("OK-NOTE", 0), cnt.get("NOT OK", 0), cnt.get("UNVERIFIABLE", 0), sum(cnt.values())])
    ss.append([])
    ss.append(["Legend"])
    ss.append(["OK", "Reviewer's verdict and note are acceptable as written."])
    ss.append(["OK-NOTE", "Verdict acceptable; note needs the correction described, or a policy decision is needed (second-person style)."])
    ss.append(["NOT OK", "Verdict or note is wrong, unsupported, or does not answer the question asked. Must be redone before sign-off."])
    ss.append(["UNVERIFIABLE", "Cannot be checked against the sources supplied to this project; reviewer must add the LOM reference."])
    ss.append([])
    ss.append(["QA performed against the statute texts supplied to the project (versions in the Sources & currency tab). Red Book and Sign-off tabs were not reviewed, as instructed."])
    for col, w in zip("ABCDEF", (22, 8, 10, 10, 14, 14)):
        ss.column_dimensions[col].width = w
    wb.save(out)
    for tab, cnt in summary.items():
        print(tab, dict(cnt))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
